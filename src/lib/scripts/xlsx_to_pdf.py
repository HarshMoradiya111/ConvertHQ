from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak


def sheet_rows(sheet):
    max_row = sheet.max_row or 1
    max_column = sheet.max_column or 1

    for row_index in range(1, max_row + 1):
        row_values = []
        for col_index in range(1, max_column + 1):
            value = sheet.cell(row=row_index, column=col_index).value
            row_values.append("" if value is None else str(value))
        yield row_values


def convert_xlsx_to_pdf(input_xlsx: Path, output_pdf: Path) -> None:
    workbook = load_workbook(str(input_xlsx), data_only=True)
    output_pdf.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(output_pdf),
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SheetTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        spaceAfter=10,
    )

    story = []

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        if sheet_index > 1:
            story.append(PageBreak())

        story.append(Paragraph(sheet.title, title_style))
        story.append(Spacer(1, 0.15 * inch))

        rows = list(sheet_rows(sheet))
        if not rows:
            story.append(Paragraph("(Empty sheet)", styles["BodyText"]))
            continue

        table = Table(rows, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)

    doc.build(story)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: xlsx_to_pdf.py <input_xlsx> <output_pdf>", file=sys.stderr)
        return 1

    input_xlsx = Path(sys.argv[1])
    output_pdf = Path(sys.argv[2])

    if not input_xlsx.exists():
        print(f"Input XLSX not found: {input_xlsx}", file=sys.stderr)
        return 1

    try:
        convert_xlsx_to_pdf(input_xlsx, output_pdf)
        print(f'{{"inputXlsx":"{input_xlsx.as_posix()}","outputPdf":"{output_pdf.as_posix()}"}}')
        return 0
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())