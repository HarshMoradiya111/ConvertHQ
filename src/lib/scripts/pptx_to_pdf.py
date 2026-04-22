from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def slide_text_frames(slide):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip()
        if text:
            yield text


def table_rows_for_slide(slide):
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            rows = []
            table = shape.table
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            if rows:
                yield rows


def convert_pptx_to_pdf(input_pptx: Path, output_pdf: Path) -> None:
    presentation = Presentation(str(input_pptx))
    width = presentation.slide_width / 914400
    height = presentation.slide_height / 914400
    page_size = landscape((width * inch, height * inch)) if width >= height else portrait((width * inch, height * inch))

    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SlideTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        spaceAfter=10,
    )
    body_style = ParagraphStyle(
        "SlideBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        spaceAfter=8,
    )

    doc = SimpleDocTemplate(
        str(output_pdf),
        pagesize=page_size,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )

    story = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        if slide_index > 1:
            story.append(Spacer(1, 0.25 * inch))

        title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.has_text_frame else f"Slide {slide_index}"
        story.append(Paragraph(title, title_style))

        for text in slide_text_frames(slide):
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            if not lines:
                continue
            if len(lines) == 1:
                story.append(Paragraph(lines[0].replace("\n", "<br/>"), body_style))
            else:
                for line in lines:
                    story.append(Paragraph(f"• {line}", body_style))

        for rows in table_rows_for_slide(slide):
            table = Table(rows, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(table)

    doc.build(story)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: pptx_to_pdf.py <input_pptx> <output_pdf>", file=sys.stderr)
        return 1

    input_pptx = Path(sys.argv[1])
    output_pdf = Path(sys.argv[2])

    if not input_pptx.exists():
        print(f"Input PPTX not found: {input_pptx}", file=sys.stderr)
        return 1

    try:
        convert_pptx_to_pdf(input_pptx, output_pdf)
        print(f'{{"inputPptx":"{input_pptx.as_posix()}","outputPdf":"{output_pdf.as_posix()}"}}')
        return 0
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())