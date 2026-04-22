from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt


def configure_section(section, page_width_points: float, page_height_points: float) -> None:
    section.page_width = Pt(page_width_points)
    section.page_height = Pt(page_height_points)
    section.left_margin = Inches(0)
    section.right_margin = Inches(0)
    section.top_margin = Inches(0)
    section.bottom_margin = Inches(0)


def configure_document_style(document: Document) -> None:
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Courier New"
    normal_style.font.size = Pt(10)


def add_image_page(document: Document, pdf_page, page_width_points: float, page_height_points: float) -> None:
    bitmap = pdf_page.render(scale=2)
    image = bitmap.to_pil()
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)

    document.add_picture(
        buffer,
        width=Inches(page_width_points / 72.0),
        height=Inches(page_height_points / 72.0),
    )


def convert_pdf_to_docx(input_pdf: Path, output_docx: Path) -> None:
    document = Document()
    configure_document_style(document)

    pdfium_doc = pdfium.PdfDocument(str(input_pdf))
    try:
        for page_index in range(len(pdfium_doc)):
            pdfium_page = pdfium_doc.get_page(page_index)
            try:
                if page_index > 0:
                    document.add_section(WD_SECTION.NEW_PAGE)

                width_points = float(pdfium_doc.get_page_size(page_index)[0])
                height_points = float(pdfium_doc.get_page_size(page_index)[1])
                configure_section(document.sections[-1], width_points, height_points)

                add_image_page(document, pdfium_page, width_points, height_points)
            finally:
                pdfium_page.close()
    finally:
        pdfium_doc.close()

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_docx))


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: pdf_to_docx.py <input_pdf> <output_docx>", file=sys.stderr)
        return 1

    input_pdf = Path(sys.argv[1])
    output_docx = Path(sys.argv[2])

    if not input_pdf.exists():
        print(f"Input PDF not found: {input_pdf}", file=sys.stderr)
        return 1

    try:
        convert_pdf_to_docx(input_pdf, output_docx)
        print(
            f'{{"inputPdf":"{input_pdf.as_posix()}","outputDocx":"{output_docx.as_posix()}"}}'
        )
        return 0
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())